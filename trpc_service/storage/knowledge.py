"""Tenant-owned knowledge ingestion, retrieval, and file normalization."""

import asyncio
import json
import sys
from zipfile import ZipFile
import hashlib
import logging
from collections.abc import AsyncIterator, Mapping, Sequence
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from io import BytesIO
from pathlib import Path
from typing import BinaryIO
from uuid import UUID, uuid4

from docx import Document as DocxDocument
from pypdf import PdfReader
from sqlalchemy import delete, func, select, text as sql_text
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker

from trpc_service.storage.knowledge_orm import (
    KnowledgeArtifactRow,
    KnowledgeBaseRow,
    KnowledgeChunkRow,
    KnowledgeDocumentRow,
)
from trpc_service.storage.ports import ArtifactStore, KnowledgeStore
from trpc_service.storage.router import BackendProfile, StorageRouter
from trpc_service.storage.types import (
    ArtifactMetadata,
    ArtifactRef,
    KnowledgeDocument,
    KnowledgeHit,
)
from trpc_service.tenant.context import TenantContext

MAX_KNOWLEDGE_FILE_BYTES = 10 * 1024 * 1024
logger = logging.getLogger(__name__)


@dataclass(frozen=True, slots=True)
class TextChunk:
    """One deterministic source range ready for embedding and citation."""

    index: int
    content: str
    start_char: int
    end_char: int


class KnowledgeFileParser:
    """Extract bounded textual content from the supported knowledge formats."""

    _TEXT_SUFFIXES = frozenset({".txt", ".md", ".markdown"})
    _SUPPORTED_SUFFIXES = _TEXT_SUFFIXES | frozenset({".pdf", ".docx"})

    def __init__(self, *, max_extracted_chars: int = 2_000_000) -> None:
        if max_extracted_chars < 1:
            raise ValueError("extracted text limit must be positive")
        self._max_extracted_chars = max_extracted_chars

    def supports(self, filename: str) -> bool:
        """Return whether the filename selects an explicitly supported parser."""

        return Path(filename).suffix.lower() in self._SUPPORTED_SUFFIXES

    def parse(self, filename: str, media_type: str, source: BinaryIO) -> str:
        """Return normalized text while rejecting unapproved file formats."""

        suffix = Path(filename).suffix.lower()
        source.seek(0)
        if suffix in self._TEXT_SUFFIXES or media_type in {"text/plain", "text/markdown"}:
            try:
                text = source.read().decode("utf-8-sig")
            except UnicodeDecodeError as error:
                raise ValueError("knowledge text file must use UTF-8 encoding") from error
        elif suffix == ".pdf" or media_type == "application/pdf":
            reader = PdfReader(source)
            if len(reader.pages) > 500:
                raise ValueError("knowledge PDF exceeds the 500 page limit")
            parts = []
            extracted = 0
            for page in reader.pages:
                part = page.extract_text() or ""
                extracted += len(part)
                if extracted > self._max_extracted_chars:
                    raise ValueError("knowledge file has too much extracted text")
                parts.append(part)
            text = "\n\n".join(parts)
        elif suffix == ".docx" or media_type == (
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document"):
            # python-docx accepts a seekable file-like object; copying avoids a
            # provider-specific stream retaining ownership of the upload handle.
            with ZipFile(source) as archive:
                if (len(archive.infolist()) > 2000
                        or sum(item.file_size for item in archive.infolist()) > 32 * 1024 * 1024):
                    raise ValueError("knowledge DOCX exceeds the expanded file limit")
            source.seek(0)
            document = DocxDocument(BytesIO(source.read()))
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        else:
            raise ValueError("unsupported knowledge file; use TXT, Markdown, PDF, or DOCX")

        normalized = text.replace("\r\n", "\n").replace("\r", "\n").strip()
        if not normalized:
            raise ValueError("knowledge file contains no extractable text")
        if len(normalized) > self._max_extracted_chars:
            raise ValueError("knowledge file has too much extracted text")
        return normalized


class TextChunker:
    """Split source text deterministically with overlap for retrieval continuity."""

    def __init__(self, *, chunk_size: int = 1_200, overlap: int = 200) -> None:
        if chunk_size < 1 or overlap < 0 or overlap >= chunk_size:
            raise ValueError("chunk size must be positive and overlap smaller than it")
        self._chunk_size = chunk_size
        self._overlap = overlap

    def split(self, text: str) -> tuple[TextChunk, ...]:
        """Return non-empty chunks with offsets into the normalized source text."""

        normalized = text.strip()
        if not normalized:
            return ()
        chunks: list[TextChunk] = []
        start = 0
        while start < len(normalized):
            hard_end = min(start + self._chunk_size, len(normalized))
            end = hard_end
            if hard_end < len(normalized):
                # Prefer a paragraph/line/space boundary without producing a
                # pathologically small chunk for text that has no separators.
                minimum = start + self._chunk_size // 2
                candidates = [
                    normalized.rfind(separator, minimum, hard_end)
                    for separator in ("\n\n", "\n", "。", ". ", " ")
                ]
                boundary = max(candidates)
                if boundary >= minimum:
                    end = boundary + 1
            content = normalized[start:end].strip()
            if content:
                chunks.append(
                    TextChunk(
                        index=len(chunks),
                        content=content,
                        start_char=start,
                        end_char=end,
                    ))
            if end >= len(normalized):
                break
            start = max(start + 1, end - self._overlap)
        return tuple(chunks)


@dataclass(frozen=True, slots=True)
class KnowledgeArtifact:
    """Tenant-owned uploaded object visible to the ingestion service."""

    artifact_id: str
    filename: str
    media_type: str
    checksum: str
    size_bytes: int


@dataclass(frozen=True, slots=True)
class KnowledgeDocumentState:
    """Stable public state of one versioned Knowledge Document."""

    document_id: UUID
    knowledge_base_id: UUID
    filename: str
    version: int
    status: str
    chunk_count: int


def _allowed_base_names(config: Mapping[str, object]) -> frozenset[str]:
    """Validate fail-closed Agent knowledge access configuration."""

    raw = config.get("knowledge_base_names", ())
    if not isinstance(raw, Sequence) or isinstance(raw, (str, bytes)):
        raise ValueError("knowledge_base_names must be an array")
    if any(not isinstance(name, str) or not name.strip() for name in raw):
        raise ValueError("knowledge_base_names must contain non-empty strings")
    return frozenset(str(name).strip() for name in raw)


def _require_base_access(config: Mapping[str, object], name: str) -> None:
    """Keep Agent configuration, rather than model arguments, authoritative."""

    if name not in _allowed_base_names(config):
        raise PermissionError("knowledge base is not granted to this Agent")


def _document_state(row: KnowledgeDocumentRow) -> KnowledgeDocumentState:
    return KnowledgeDocumentState(
        document_id=row.document_id,
        knowledge_base_id=row.knowledge_base_id,
        filename=row.filename,
        version=row.version,
        status=row.status,
        chunk_count=row.chunk_count,
    )


async def _lock_scope(database: AsyncSession, key: str) -> None:
    """Serialize idempotency decisions across PostgreSQL Worker processes."""

    bind = database.get_bind()
    if bind.dialect.name == "postgresql":
        await database.execute(
            sql_text("SELECT pg_advisory_xact_lock(hashtextextended(:key, 0))"),
            {"key": key},
        )


class TenantKnowledgeService:
    """Coordinate SQL metadata, object storage, and a replaceable vector index."""

    def __init__(
        self,
        sessions: async_sessionmaker[AsyncSession],
        *,
        knowledge: KnowledgeStore | None = None,
        artifacts: ArtifactStore | None = None,
        storage: StorageRouter | None = None,
        default_backends: Mapping[str, object] | None = None,
        embedding_model: str = "text-embedding-v4",
        max_file_bytes: int = MAX_KNOWLEDGE_FILE_BYTES,
        ingest_lease_seconds: int = 15 * 60,
    ) -> None:
        if max_file_bytes < 1 or ingest_lease_seconds < 1:
            raise ValueError("knowledge file limit and ingestion lease must be positive")
        if storage is None and (knowledge is None or artifacts is None):
            raise ValueError("knowledge service requires fixed stores or a storage router")
        if storage is not None and (knowledge is not None or artifacts is not None):
            raise ValueError("knowledge service cannot mix fixed stores and a storage router")
        if not embedding_model.strip():
            raise ValueError("embedding model must not be empty")
        self._embedding_model = embedding_model
        self._sessions = sessions
        self._knowledge = knowledge
        self._artifacts = artifacts
        self._storage = storage
        self._default_backends = dict(default_backends or {})
        self._parser = KnowledgeFileParser()
        self._parse_slots = asyncio.Semaphore(2)
        self._chunker = TextChunker()
        self._max_file_bytes = max_file_bytes
        self._ingest_lease_seconds = ingest_lease_seconds

    async def _parse(self, filename: str, media_type: str, payload: bytes) -> str:
        async with self._parse_slots:
            if Path(filename).suffix.lower() in self._parser._TEXT_SUFFIXES:
                return await asyncio.to_thread(self._parser.parse, filename, media_type,
                                               BytesIO(payload))
            process = await asyncio.create_subprocess_exec(
                sys.executable,
                "-m",
                "trpc_service.storage.parse_worker",
                filename,
                media_type,
                stdin=asyncio.subprocess.PIPE,
                stdout=asyncio.subprocess.PIPE,
                stderr=asyncio.subprocess.DEVNULL,
            )
            try:
                async with asyncio.timeout(25):
                    stdout, _ = await process.communicate(payload)
                if process.returncode != 0:
                    raise ValueError("document parsing failed or exceeded its resource limit")
                result = json.loads(stdout)
                if not isinstance(result, str):
                    raise ValueError("invalid document parser response")
                return result
            finally:
                if process.returncode is None:
                    process.kill()
                    await process.wait()

    def _stores(
        self,
        backends: Mapping[str, object] | None,
    ) -> tuple[KnowledgeStore, ArtifactStore]:
        """Resolve the immutable Agent Backend Profile for one operation."""

        if self._storage is None:
            if self._knowledge is None or self._artifacts is None:
                raise RuntimeError("fixed knowledge storage is incomplete")
            return self._knowledge, self._artifacts
        profile = backends or self._default_backends
        resolved = self._storage.resolve(BackendProfile.from_mapping(profile))
        if resolved.knowledge is None or resolved.artifact is None:
            raise ValueError("Agent backend profile requires knowledge and artifact storage")
        return resolved.knowledge, resolved.artifact

    def supports_upload(self, filename: str) -> bool:
        """Return whether a tenant-console upload can become a knowledge source."""

        return self._parser.supports(filename)

    async def register_artifact(
        self,
        context: TenantContext,
        *,
        principal_id: str,
        reference: ArtifactRef,
        metadata: ArtifactMetadata,
    ) -> KnowledgeArtifact:
        """Register an already persisted tenant object as a knowledge candidate."""

        if not principal_id.strip() or not metadata.filename.strip(
        ) or not metadata.media_type.strip():
            raise ValueError("knowledge upload identity and file metadata are required")
        if not self._parser.supports(metadata.filename):
            raise ValueError("unsupported knowledge file; use TXT, Markdown, PDF, or DOCX")
        if metadata.size_bytes > self._max_file_bytes:
            raise ValueError("knowledge file exceeds the 10 MB limit")
        if reference.checksum != metadata.checksum:
            raise ValueError("knowledge artifact checksum does not match its object reference")
        async with self._sessions.begin() as database:
            await _lock_scope(database, f"artifact:{context.tenant_id}:{metadata.checksum}")
            row = await database.get(
                KnowledgeArtifactRow,
                (context.tenant_id, reference.artifact_id),
            )
            if row is None:
                database.add(
                    KnowledgeArtifactRow(
                        tenant_id=context.tenant_id,
                        artifact_id=reference.artifact_id,
                        uploaded_by_agent_id=context.agent_app_id,
                        uploaded_by_principal_id=principal_id,
                        filename=metadata.filename,
                        media_type=metadata.media_type,
                        checksum=metadata.checksum,
                        size_bytes=metadata.size_bytes,
                        object_uri=reference.uri,
                    ))
            elif row.checksum != metadata.checksum:
                raise RuntimeError("artifact identifier collision")
        return KnowledgeArtifact(
            reference.artifact_id,
            metadata.filename,
            metadata.media_type,
            metadata.checksum,
            metadata.size_bytes,
        )

    async def upload(
        self,
        context: TenantContext,
        *,
        principal_id: str,
        filename: str,
        media_type: str,
        content: AsyncIterator[bytes],
        backends: Mapping[str, object] | None = None,
    ) -> KnowledgeArtifact:
        """Validate one bounded upload before making its object and metadata visible."""

        if not principal_id.strip() or not filename.strip() or not media_type.strip():
            raise ValueError("knowledge upload identity and file metadata are required")
        if not self._parser.supports(filename):
            raise ValueError("unsupported knowledge file; use TXT, Markdown, PDF, or DOCX")
        payload = bytearray()
        async for block in content:
            payload.extend(block)
            if len(payload) > self._max_file_bytes:
                raise ValueError("knowledge file exceeds the 10 MB limit")
        checksum = hashlib.sha256(payload).hexdigest()
        metadata = ArtifactMetadata(
            filename=filename,
            media_type=media_type,
            checksum=checksum,
            size_bytes=len(payload),
        )

        async def blocks() -> AsyncIterator[bytes]:
            yield bytes(payload)

        _, artifacts = self._stores(backends)
        reference = await artifacts.put(context, blocks(), metadata)
        return await self.register_artifact(
            context,
            principal_id=principal_id,
            reference=reference,
            metadata=metadata,
        )

    async def ingest(
        self,
        context: TenantContext,
        config: Mapping[str, object],
        *,
        knowledge_base_name: str,
        artifact_ids: Sequence[str],
        backends: Mapping[str, object] | None = None,
    ) -> tuple[KnowledgeDocumentState, ...]:
        """Idempotently ingest trusted message attachments into one granted base."""

        _require_base_access(config, knowledge_base_name)
        if not artifact_ids:
            raise ValueError("knowledge ingestion requires at least one uploaded artifact")
        results: list[KnowledgeDocumentState] = []
        for artifact_id in dict.fromkeys(artifact_ids):
            results.append(await self._ingest_one(
                context,
                knowledge_base_name,
                artifact_id,
                backends,
            ))
        return tuple(results)

    async def update_document(
        self,
        context: TenantContext,
        config: Mapping[str, object],
        knowledge_base_name: str,
        document_id: UUID,
        artifact_id: str,
        *,
        backends: Mapping[str, object] | None = None,
    ) -> KnowledgeDocumentState:
        """Create the next version of one tenant document from an approved upload."""

        _require_base_access(config, knowledge_base_name)
        if not artifact_id.strip():
            raise ValueError("knowledge update requires one uploaded artifact")
        return await self._ingest_one(
            context,
            knowledge_base_name,
            artifact_id,
            backends,
            replace_document_id=document_id,
        )

    async def _ingest_one(
        self,
        context: TenantContext,
        knowledge_base_name: str,
        artifact_id: str,
        backends: Mapping[str, object] | None,
        *,
        replace_document_id: UUID | None = None,
    ) -> KnowledgeDocumentState:
        """Run one recoverable ingestion lifecycle with durable status transitions."""

        knowledge, artifacts = self._stores(backends)
        async with self._sessions.begin() as database:
            await _lock_scope(
                database,
                f"knowledge:{context.tenant_id}:{knowledge_base_name}",
            )
            artifact = await database.get(
                KnowledgeArtifactRow,
                (context.tenant_id, artifact_id),
            )
            if artifact is None:
                raise LookupError("knowledge artifact does not exist in this tenant")
            base = await database.scalar(
                select(KnowledgeBaseRow).where(
                    KnowledgeBaseRow.tenant_id == context.tenant_id,
                    KnowledgeBaseRow.name == knowledge_base_name,
                    KnowledgeBaseRow.status == "ACTIVE",
                ))
            if base is not None:
                self._validate_embedding_model(base)
            if base is None:
                if replace_document_id is not None:
                    raise LookupError("knowledge document does not exist in this tenant base")
                base = KnowledgeBaseRow(
                    tenant_id=context.tenant_id,
                    name=knowledge_base_name,
                    embedding_model=self._embedding_model,
                )
                database.add(base)
                await database.flush()
            replaced: KnowledgeDocumentRow | None = None
            if replace_document_id is not None:
                replaced = await database.scalar(
                    select(KnowledgeDocumentRow).where(
                        KnowledgeDocumentRow.tenant_id == context.tenant_id,
                        KnowledgeDocumentRow.knowledge_base_id == base.knowledge_base_id,
                        KnowledgeDocumentRow.document_id == replace_document_id,
                        KnowledgeDocumentRow.status == "READY",
                    ))
                if replaced is None:
                    raise LookupError("knowledge document does not exist in this tenant base")
            duplicate = await database.scalar(
                select(KnowledgeDocumentRow).where(
                    KnowledgeDocumentRow.tenant_id == context.tenant_id,
                    KnowledgeDocumentRow.knowledge_base_id == base.knowledge_base_id,
                    KnowledgeDocumentRow.checksum == artifact.checksum,
                ))
            if duplicate is not None and duplicate.status == "READY":
                if replaced is None or duplicate.document_id == replaced.document_id:
                    return _document_state(duplicate)
                raise ValueError("replacement content already exists in this knowledge base")
            if duplicate is not None and duplicate.status == "INGESTING":
                updated_at = duplicate.updated_at
                if updated_at.tzinfo is None:
                    updated_at = updated_at.replace(tzinfo=timezone.utc)
                if updated_at > datetime.now(
                        timezone.utc) - timedelta(seconds=self._ingest_lease_seconds):
                    # Another Worker owns the same checksum. Returning the
                    # durable state avoids duplicate indexing across nodes.
                    return _document_state(duplicate)
            logical_filename = replaced.filename if replaced is not None else artifact.filename
            reactivating_deleted = duplicate is not None and duplicate.status == "DELETED"
            if duplicate is None or reactivating_deleted:
                current_version = await database.scalar(
                    select(func.max(KnowledgeDocumentRow.version)).where(
                        KnowledgeDocumentRow.tenant_id == context.tenant_id,
                        KnowledgeDocumentRow.knowledge_base_id == base.knowledge_base_id,
                        KnowledgeDocumentRow.filename == logical_filename,
                    ))
                version = int(current_version or 0) + 1
            else:
                version = duplicate.version
            document = duplicate or KnowledgeDocumentRow(
                tenant_id=context.tenant_id,
                document_id=uuid4(),
                knowledge_base_id=base.knowledge_base_id,
                artifact_id=artifact.artifact_id,
                filename=logical_filename,
                checksum=artifact.checksum,
                version=version,
                status="INGESTING",
            )
            retry_chunk_ids: list[str] = []
            if duplicate is None:
                database.add(document)
            else:
                # A superseded checksum cannot become the latest version by an
                # accidental duplicate upload. An explicitly re-added deleted
                # source is different: rebuild its projection and advance its
                # version while retaining the stable document identity.
                if duplicate.status == "SUPERSEDED":
                    raise ValueError(
                        "a tombstoned checksum cannot be re-ingested; upload revised content")
                if reactivating_deleted:
                    document.filename = logical_filename
                    document.version = version
                    document.deleted_at = None
                document.status = "INGESTING"
                document.error_summary = None
                document.chunk_count = 0
                document.updated_at = datetime.now(timezone.utc)
                retry_chunk_ids = list((await database.scalars(
                    select(KnowledgeChunkRow.chunk_id).where(
                        KnowledgeChunkRow.tenant_id == context.tenant_id,
                        KnowledgeChunkRow.document_id == document.document_id,
                    ))).all())
                await database.execute(
                    delete(KnowledgeChunkRow).where(
                        KnowledgeChunkRow.tenant_id == context.tenant_id,
                        KnowledgeChunkRow.document_id == document.document_id,
                    ))
            await database.flush()
            document_id = document.document_id
            base_id = base.knowledge_base_id
            filename = document.filename
            media_type = artifact.media_type

        vector_document_ids: list[str] = []
        try:
            if retry_chunk_ids:
                await knowledge.delete(context, str(base_id), retry_chunk_ids)
            payload = b"".join([block async for block in artifacts.open(context, artifact_id)])
            text = await self._parse(filename, media_type, payload)
            chunks = self._chunker.split(text)
            if not chunks:
                raise ValueError("knowledge file produced no chunks")
            vector_documents = [
                KnowledgeDocument(
                    document_id=f"{document_id}:{chunk.index}",
                    knowledge_base_id=str(base_id),
                    content=chunk.content,
                    attributes={
                        "source_document_id": str(document_id),
                        "filename": filename,
                        "version": version,
                        "chunk_index": chunk.index,
                        "start_char": chunk.start_char,
                        "end_char": chunk.end_char,
                    },
                ) for chunk in chunks
            ]
            vector_document_ids = [item.document_id for item in vector_documents]
            await knowledge.index(context, vector_documents)
            async with self._sessions.begin() as database:
                previous = (await database.scalars(
                    select(KnowledgeDocumentRow).where(
                        KnowledgeDocumentRow.tenant_id == context.tenant_id,
                        KnowledgeDocumentRow.knowledge_base_id == base_id,
                        KnowledgeDocumentRow.filename == filename,
                        KnowledgeDocumentRow.status == "READY",
                        KnowledgeDocumentRow.document_id != document_id,
                    ))).all()
                old_chunk_ids: list[str] = []
                for prior in previous:
                    prior.status = "SUPERSEDED"
                    prior.deleted_at = datetime.now(timezone.utc)
                    old_chunk_ids.extend((await database.scalars(
                        select(KnowledgeChunkRow.chunk_id).where(
                            KnowledgeChunkRow.tenant_id == context.tenant_id,
                            KnowledgeChunkRow.document_id == prior.document_id,
                        ))).all())
                stored_document = await database.get(
                    KnowledgeDocumentRow,
                    (context.tenant_id, document_id),
                )
                if stored_document is None:
                    raise RuntimeError("knowledge ingestion state disappeared")
                for chunk, vector_document in zip(chunks, vector_documents, strict=True):
                    database.add(
                        KnowledgeChunkRow(
                            tenant_id=context.tenant_id,
                            chunk_id=vector_document.document_id,
                            document_id=document_id,
                            chunk_index=chunk.index,
                            start_char=chunk.start_char,
                            end_char=chunk.end_char,
                            content_hash=hashlib.sha256(chunk.content.encode()).hexdigest(),
                            content=chunk.content,
                            attributes=dict(vector_document.attributes),
                        ))
                stored_document.status = "READY"
                stored_document.chunk_count = len(chunks)
                stored_document.error_summary = None
        except Exception as error:
            # Projection writes precede the authoritative SQL transition. A
            # failed transition must make best-effort removal before allowing a
            # retry to reuse the same deterministic chunk identifiers.
            if vector_document_ids:
                try:
                    await knowledge.delete(context, str(base_id), vector_document_ids)
                except Exception as cleanup_error:
                    logger.warning(
                        "Failed to remove partial knowledge vectors error_type=%s",
                        type(cleanup_error).__name__,
                    )
            async with self._sessions.begin() as database:
                failed = await database.get(
                    KnowledgeDocumentRow,
                    (context.tenant_id, document_id),
                )
                if failed is not None:
                    failed.status = "FAILED"
                    failed.error_summary = type(error).__name__
            raise
        # Old projections cannot be returned because search authorizes READY
        # SQL rows. Cleanup is best effort so a vector outage cannot turn an
        # already committed READY document back into FAILED.
        try:
            await knowledge.delete(context, str(base_id), old_chunk_ids)
        except Exception as cleanup_error:
            logger.warning(
                "Deferred superseded-vector cleanup failed error_type=%s",
                type(cleanup_error).__name__,
            )
        return KnowledgeDocumentState(document_id, base_id, filename, version, "READY", len(chunks))

    async def list_documents(
        self,
        context: TenantContext,
        config: Mapping[str, object],
        knowledge_base_name: str,
        backends: Mapping[str, object] | None = None,
    ) -> tuple[KnowledgeDocumentState, ...]:
        """List only active documents in one explicitly granted tenant base."""

        _require_base_access(config, knowledge_base_name)
        async with self._sessions() as database:
            base = await database.scalar(
                select(KnowledgeBaseRow).where(
                    KnowledgeBaseRow.tenant_id == context.tenant_id,
                    KnowledgeBaseRow.name == knowledge_base_name,
                    KnowledgeBaseRow.status == "ACTIVE",
                ))
            if base is None:
                return ()
            rows = (await database.scalars(
                select(KnowledgeDocumentRow).where(
                    KnowledgeDocumentRow.tenant_id == context.tenant_id,
                    KnowledgeDocumentRow.knowledge_base_id == base.knowledge_base_id,
                    KnowledgeDocumentRow.status == "READY",
                ).order_by(KnowledgeDocumentRow.filename,
                           KnowledgeDocumentRow.version.desc()))).all()
        return tuple(_document_state(row) for row in rows)

    def _validate_embedding_model(self, base: KnowledgeBaseRow) -> None:
        if base.embedding_model != self._embedding_model:
            raise ValueError("knowledge embedding model changed; reindex this knowledge base first")

    async def search(
        self,
        context: TenantContext,
        config: Mapping[str, object],
        query: str,
        *,
        limit: int = 5,
        backends: Mapping[str, object] | None = None,
    ) -> tuple[KnowledgeHit, ...]:
        """Search all granted tenant bases and merge results by score."""

        if not query.strip() or limit < 1 or limit > 50:
            raise ValueError("knowledge query and a limit from 1 to 50 are required")
        knowledge, _ = self._stores(backends)
        names = _allowed_base_names(config)
        if not names:
            return ()
        async with self._sessions() as database:
            bases = (await database.scalars(
                select(KnowledgeBaseRow).where(
                    KnowledgeBaseRow.tenant_id == context.tenant_id,
                    KnowledgeBaseRow.name.in_(names),
                    KnowledgeBaseRow.status == "ACTIVE",
                ))).all()
        for base in bases:
            self._validate_embedding_model(base)
        base_ids = [base.knowledge_base_id for base in bases]
        hits: list[KnowledgeHit] = []
        for base_id in base_ids:
            hits.extend(await knowledge.search(context, str(base_id), query, limit))
        scoped_hits: list[tuple[KnowledgeHit, UUID]] = []
        for hit in hits:
            raw_source_id = hit.document.attributes.get("source_document_id")
            try:
                source_id = UUID(str(raw_source_id))
            except (TypeError, ValueError):
                # Vector rows are projections, not authority. A row without a
                # valid SQL source reference is never exposed to the Runner.
                continue
            scoped_hits.append((hit, source_id))
        if not scoped_hits:
            return ()
        source_ids = {source_id for _, source_id in scoped_hits}
        async with self._sessions() as database:
            ready_ids = frozenset((await database.scalars(
                select(KnowledgeDocumentRow.document_id).where(
                    KnowledgeDocumentRow.tenant_id == context.tenant_id,
                    KnowledgeDocumentRow.knowledge_base_id.in_(base_ids),
                    KnowledgeDocumentRow.document_id.in_(source_ids),
                    KnowledgeDocumentRow.status == "READY",
                ))).all())
        authoritative = [hit for hit, source_id in scoped_hits if source_id in ready_ids]
        return tuple(sorted(authoritative, key=lambda hit: -hit.score)[:limit])

    async def delete_document(
        self,
        context: TenantContext,
        config: Mapping[str, object],
        knowledge_base_name: str,
        document_id: UUID,
        backends: Mapping[str, object] | None = None,
    ) -> KnowledgeDocumentState:
        """Idempotently soft-delete one active or already deleted tenant document."""

        _require_base_access(config, knowledge_base_name)
        knowledge, _ = self._stores(backends)
        async with self._sessions() as database:
            document = await database.scalar(
                select(KnowledgeDocumentRow).join(
                    KnowledgeBaseRow,
                    (KnowledgeBaseRow.tenant_id == KnowledgeDocumentRow.tenant_id)
                    &
                    (KnowledgeBaseRow.knowledge_base_id == KnowledgeDocumentRow.knowledge_base_id),
                ).where(
                    KnowledgeDocumentRow.tenant_id == context.tenant_id,
                    KnowledgeDocumentRow.document_id == document_id,
                    KnowledgeBaseRow.name == knowledge_base_name,
                    KnowledgeBaseRow.status == "ACTIVE",
                ))
            if document is None:
                raise LookupError("knowledge document does not exist in this tenant base")
            if document.status == "DELETED":
                # IM conversations may replay a desired-state delete from model
                # history. A tenant-scoped tombstone proves that it already won.
                return _document_state(document)
            if document.status != "READY":
                raise LookupError("knowledge document does not exist in this tenant base")
            chunk_ids = (await database.scalars(
                select(KnowledgeChunkRow.chunk_id).where(
                    KnowledgeChunkRow.tenant_id == context.tenant_id,
                    KnowledgeChunkRow.document_id == document_id,
                ))).all()
            base_id = document.knowledge_base_id
        # Remove retrieval visibility first. If the metadata transition fails,
        # an operator can retry deletion without leaking a tombstoned source.
        await knowledge.delete(context, str(base_id), chunk_ids)
        async with self._sessions.begin() as database:
            document = await database.get(
                KnowledgeDocumentRow,
                (context.tenant_id, document_id),
            )
            if document is None:
                raise RuntimeError("knowledge document changed during deletion")
            if document.status == "DELETED":
                # Concurrent confirmations can both remove the same vector
                # projection; the committed tombstone is a successful result.
                return _document_state(document)
            if document.status != "READY":
                raise RuntimeError("knowledge document changed during deletion")
            document.status = "DELETED"
            document.deleted_at = datetime.now(timezone.utc)
            state = _document_state(document)
        return KnowledgeDocumentState(
            state.document_id,
            state.knowledge_base_id,
            state.filename,
            state.version,
            "DELETED",
            state.chunk_count,
        )
